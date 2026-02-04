import openpyxl

wb = openpyxl.load_workbook('test_output_payload.xlsx')
print('シート一覧:', wb.sheetnames)

for sheet_name in wb.sheetnames:
    print(f"\n=== シート: {sheet_name} ===")
    ws = wb[sheet_name]
    
    for idx, row in enumerate(list(ws.rows)[:12], 1):
        print(f"{idx}: {[cell.value for cell in row]}")
