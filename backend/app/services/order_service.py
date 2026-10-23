"""発注書ペイロード生成 + メール/PDF/Excel レンダラー。

承認済み PurchaseRequest を業者別にグループ化し、
- 正本 JSON (payload)
- メール本文 (テキスト)
- PDF (fpdf2)
- Excel (openpyxl)
を生成する。
"""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings
from app.models import Item, PurchaseRequest, User
from app.utils.datetime import jst_today
from app.utils.xlsx import write_cell

FONT_NAME = settings.font_name
FONT_PATH: Path = settings.font_path
# 発注書を出せる状態（発注済からの出し直しも可：C-06）
ORDERABLE_STATUSES = ("APPROVED", "ORDERED")


def _format_currency(yen: float | None) -> str:
    return "要確認" if yen is None else f"¥{yen:,.0f}"


def build_payloads(db: Session, request_ids: list[int]) -> list[dict[str, Any]]:
    """承認済み PR から業者別 payload を生成。"""
    if not request_ids:
        return []
    rows = db.execute(
        select(PurchaseRequest, Item, User)
        .join(Item, PurchaseRequest.item_id == Item.id)
        .outerjoin(User, PurchaseRequest.requested_by_user_id == User.id)
        .where(PurchaseRequest.id.in_(request_ids))
        .where(PurchaseRequest.status.in_(ORDERABLE_STATUSES))
        .where(PurchaseRequest.deleted_at.is_(None))
        .order_by(Item.vendor_name, Item.item_name)
    ).all()

    vendor_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for pr, item, user in rows:
        vendor = item.vendor_name or "不明業者"
        unit_price = float(item.unit_price) if item.unit_price else None
        subtotal = unit_price * pr.requested_qty if unit_price else None
        vendor_groups[vendor].append(
            {
                "request_id": pr.id,
                "item_name": item.item_name or "品目名不明",
                "item_code": item.item_code or "",
                "vendor_part_no": item.vendor_item_code or "",
                "quantity": pr.requested_qty,
                "unit_price": unit_price,
                "subtotal": subtotal,
                "notes": pr.note or "",
                "department": pr.requested_by_department or "",
                "requester_name": user.display_name if user else "",
            }
        )

    ordered_at = jst_today().isoformat()
    payloads: list[dict[str, Any]] = []
    for vendor, items in vendor_groups.items():
        department = items[0]["department"]
        requester_name = items[0]["requester_name"]
        has_unpriced = any(i["subtotal"] is None for i in items)
        total = None if has_unpriced else sum(i["subtotal"] or 0.0 for i in items)
        payloads.append(
            {
                "vendor": vendor,
                "subject": "発注書",
                "to": f"{vendor} 御中",
                "ordered_at": ordered_at,
                "requester": {"department": department, "name": requester_name},
                "items": [
                    {k: i[k] for k in ("item_name", "item_code", "vendor_part_no", "quantity", "unit_price", "subtotal", "notes")}
                    for i in items
                ],
                "total": total,
                "has_unpriced_items": has_unpriced,
            }
        )
    return payloads


def render_email_text(payload: dict[str, Any]) -> str:
    vendor = payload.get("vendor", "")
    items = payload.get("items", [])
    ordered_at = payload.get("ordered_at", "")
    requester = payload.get("requester", {})
    department = requester.get("department", "")
    requester_name = requester.get("name", "")
    has_unpriced = payload.get("has_unpriced_items", False)
    total = payload.get("total")
    try:
        today = datetime.strptime(ordered_at, "%Y-%m-%d").strftime("%Y年%m月%d日")
    except ValueError:
        today = jst_today().strftime("%Y年%m月%d日")

    lines = []
    for idx, it in enumerate(items, 1):
        name = it.get("item_name", "")
        code = it.get("item_code", "")
        qty = it.get("quantity", 0)
        up = it.get("unit_price")
        sub = it.get("subtotal")
        if up is not None and sub is not None:
            lines.append(f"{idx}. {name}（{code}）× {qty}個 @ {_format_currency(up)} = {_format_currency(sub)}")
        else:
            lines.append(f"{idx}. {name}（{code}）× {qty}個 @ 単価未設定 = 要確認")
    total_text = _format_currency(total) + (" ※一部単価要確認" if has_unpriced and total is not None else "")
    requester_info = f"{department} {requester_name}".strip() or "担当者"
    return (
        f"件名: 【発注】物品発注のお願い\n\n"
        f"{vendor} 御中\n\n"
        f"お世話になっております。\n{requester_info}です。\n\n"
        "下記の通り発注いたしますので、納品のほどよろしくお願い申し上げます。\n\n"
        "━━━━━━━━━━━━━━━━━━━━━━━━━━\n"
        f"■ 発注日: {today}\n■ 発注明細:\n\n"
        + "\n".join(lines)
        + f"\n\n■ 合計金額: {total_text}\n"
        "━━━━━━━━━━━━━━━━━━━━━━━━━━\n\n"
        f"【納品先】\n{department or '（部署名）'}\n\n"
        f"【お問い合わせ先】\n{requester_info}\n\n"
        "以上、よろしくお願いいたします。\n"
    )


def render_pdf(payload: dict[str, Any]) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    def line(w: float, h: float, txt: str, border: int = 0, align: str = "L", *, newline: bool = False) -> None:
        nx, ny = (XPos.LMARGIN, YPos.NEXT) if newline else (XPos.RIGHT, YPos.TOP)
        pdf.cell(w, h, txt, border=border, align=align, new_x=nx, new_y=ny)

    if not FONT_PATH.exists():
        raise FileNotFoundError(f"日本語フォントが見つかりません: {FONT_PATH}")

    pdf = FPDF()
    pdf.add_font(FONT_NAME, "", str(FONT_PATH))
    pdf.add_page()
    pdf.set_font(FONT_NAME, "", 16)
    line(0, 10, "発注書", align="C", newline=True)
    pdf.ln(3)

    vendor = payload.get("vendor", "")
    items = payload.get("items", [])
    ordered_at = payload.get("ordered_at", "")
    requester = payload.get("requester", {})
    total = payload.get("total")

    pdf.set_font(FONT_NAME, "", 10)
    line(0, 6, f"発注日: {ordered_at}", newline=True)
    line(0, 6, f"発注先: {vendor}", newline=True)
    line(0, 6, f"発注元: {requester.get('department', '')} / {requester.get('name', '')}", newline=True)
    pdf.ln(5)

    pdf.set_font(FONT_NAME, "", 9)
    line(85, 8, "品目名（コード）", 1, "C")
    line(25, 8, "数量", 1, "C")
    line(35, 8, "単価", 1, "C")
    line(35, 8, "金額", 1, "C", newline=True)

    pdf.set_font(FONT_NAME, "", 8)
    for it in items:
        name = it.get("item_name", "")
        code = it.get("item_code", "")
        display = (f"{name} ({code})" if code else name)[:40]
        qty = it.get("quantity", 0)
        up = it.get("unit_price")
        sub = it.get("subtotal")
        line(85, 8, display, 1)
        line(25, 8, str(qty), 1, "C")
        if up is not None and sub is not None:
            line(35, 8, f"¥ {up:,.0f}", 1, "R")
            line(35, 8, f"¥ {sub:,.0f}", 1, "R", newline=True)
        else:
            line(35, 8, "要見積", 1, "C")
            line(35, 8, "-", 1, "C", newline=True)

    pdf.set_font(FONT_NAME, "", 9)
    line(110, 8, "")
    line(35, 8, "合計:", 1, "C")
    line(35, 8, "要見積" if total is None else f"¥ {total:,.0f}", 1, "R" if total is not None else "C", newline=True)
    pdf.ln(2)
    pdf.set_font(FONT_NAME, "", 8)
    pdf.multi_cell(0, 5, "備考: 単価未設定の品目が含まれています。" if payload.get("has_unpriced_items") else "備考: 詳細は送信メールをご確認ください。")
    return bytes(pdf.output())


def render_excel(payloads: list[dict[str, Any]]) -> bytes:
    """複数業者をシート別に出力。"""
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1565C0")
    header_font = Font(bold=True, color="FFFFFF")
    used_sheet_names: set[str] = set()
    for payload in payloads:
        vendor = payload.get("vendor", "不明業者")
        base_name = (vendor[:25] or "Unknown").replace("/", "_").replace("\\", "_").replace(":", "_").replace("*", "_").replace("?", "_").replace("[", "_").replace("]", "_")
        sheet_name = base_name
        n = 2
        while sheet_name in used_sheet_names:
            sheet_name = f"{base_name}_{n}"
            n += 1
        used_sheet_names.add(sheet_name)
        ws = wb.create_sheet(title=sheet_name[:31])

        ws["A1"] = "発注書"
        ws["A1"].font = Font(bold=True, size=16)
        ws.merge_cells("A1:G1")
        ws["A1"].alignment = Alignment(horizontal="center")

        for r, (label, value) in enumerate(
            [("発注日:", payload.get("ordered_at", "")), ("業者:", vendor),
             ("部署:", payload.get("requester", {}).get("department", "")), ("発注者:", payload.get("requester", {}).get("name", ""))],
            start=3,
        ):
            ws.cell(row=r, column=1, value=label)
            write_cell(ws, r, 2, value)

        headers = ["品目名", "品目コード", "業者品番", "数量", "単価", "金額", "備考"]
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=8, column=col, value=h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center")

        row = 9
        for it in payload.get("items", []):
            values = [
                it.get("item_name", ""), it.get("item_code", ""), it.get("vendor_part_no", ""), it.get("quantity", 0),
                it.get("unit_price") if it.get("unit_price") is not None else "未設定",
                it.get("subtotal") if it.get("subtotal") is not None else "要見積",
                it.get("notes", ""),
            ]
            for col, value in enumerate(values, 1):
                write_cell(ws, row, col, value)
            row += 1

        total = payload.get("total")
        ws.cell(row=row + 1, column=5, value="合計:").font = Font(bold=True)
        ws.cell(row=row + 1, column=6, value="要見積" if total is None else total).font = Font(bold=True)

        for col, w in enumerate([30, 14, 14, 8, 12, 14, 30], 1):
            ws.column_dimensions[chr(64 + col)].width = w

    if not wb.sheetnames:
        wb.create_sheet("Empty")
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
