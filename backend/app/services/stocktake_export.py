"""棚卸し結果 Excel レンダラー (pure openpyxl)（D-10）。

並びは保管場所→物品名（歩く順）。実数が空の行はそのまま紙の棚卸表としても使える。
"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.utils.xlsx import write_cell

COLUMNS: list[tuple[str, str, int]] = [
    ("保管場所", "location", 14),
    ("物品名", "item_name", 28),
    ("物品コード", "item_code", 14),
    ("カテゴリ", "category", 12),
    ("単位", "unit", 6),
    ("帳簿数", "system_qty", 9),
    ("実数", "actual_qty", 9),
    ("差異", "diff_qty", 9),
    ("数えた人", "counted_by", 12),
    ("数えた日時", "counted_at", 18),
    ("メモ", "note", 30),
]
HEADER_ROW = 7


def render_stocktake_excel(rows: list[dict[str, Any]], session_info: dict[str, Any], summary: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "棚卸し結果"
    header_fill = PatternFill("solid", fgColor="1565C0")
    header_font = Font(bold=True, color="FFFFFF")

    ws["A1"] = "棚卸し結果レポート"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    ws["A1"].alignment = Alignment(horizontal="center")
    for i, (label, key) in enumerate(
        [("名称", "name"), ("状態", "status"), ("開始", "started_at"), ("対象カテゴリ", "category")], start=3
    ):
        ws.cell(row=i, column=1, value=f"{label}:").font = Font(bold=True)
        write_cell(ws, i, 2, str(session_info.get(key, "")))

    for col, (label, _, width) in enumerate(COLUMNS, 1):
        c = ws.cell(row=HEADER_ROW, column=col, value=label)
        c.fill, c.font, c.alignment = header_fill, header_font, Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)

    diff_col = next(i for i, (_, k, _) in enumerate(COLUMNS, 1) if k == "diff_qty")
    for r, row in enumerate(rows, HEADER_ROW + 1):
        for col, (_, key, _) in enumerate(COLUMNS, 1):
            value = row.get(key)
            if key == "counted_at" and value:
                value = str(value).replace("T", " ")[:16]
            write_cell(ws, r, col, value)
        diff = row.get("diff_qty")
        if isinstance(diff, int) and diff:
            ws.cell(row=r, column=diff_col).font = Font(color="2E7D32" if diff > 0 else "C62828", bold=True)

    ws2 = wb.create_sheet("サマリー")
    ws2["A1"] = "サマリー"
    ws2["A1"].font = Font(bold=True, size=14)
    for i, (label, key) in enumerate(
        [("対象品目数", "total"), ("カウント済み", "counted"), ("未カウント", "uncounted"), ("一致", "matched"), ("差異あり", "diff")],
        start=3,
    ):
        ws2.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws2.cell(row=i, column=2, value=summary.get(key, 0))
    ws2.column_dimensions["A"].width = 16

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
