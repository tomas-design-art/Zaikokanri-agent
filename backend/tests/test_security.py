"""セキュリティ（受け入れ条件 G-01〜G-11）。"""
from __future__ import annotations

import io
import logging
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from app.models import Inventory, Item, User
from app.utils.datetime import jst_now

DANGEROUS = "=HYPERLINK(\"http://evil.example\",\"x\")"


def _evil_item(db) -> Item:
    it = Item(item_name=DANGEROUS, item_code="+CMD", category="消耗品", location_text="@棚", unit_price=1.0)
    db.add(it)
    db.flush()
    db.add(Inventory(item_id=it.id, quantity=3, threshold=1))
    db.commit()
    return it


# ---------- G-01 CSV / Excel の数式インジェクション ----------
def test_G01_inventory_csv_neutralizes_formulas(user_client, db):
    _evil_item(db)
    text = user_client.get("/inventory/export.csv").content.decode("utf-8-sig")
    assert f"'{DANGEROUS}".replace('"', '""') in text
    assert "'+CMD" in text
    assert "'@棚" in text


def test_G01_history_csv_neutralizes_formulas(user_client, db):
    it = _evil_item(db)
    user_client.post("/outbound", json={"item_id": it.id, "quantity": 1, "note": "-1+1"})
    text = user_client.get("/history/export.csv").content.decode("utf-8-sig")
    assert "'-1+1" in text
    assert "'=HYPERLINK" in text


def test_G01_stocktake_excel_writes_text_not_formula(admin_client, db):
    _evil_item(db)
    sid = admin_client.post("/stocktake/sessions", json={"name": "=1+1"}).json()["id"]
    ws = load_workbook(io.BytesIO(admin_client.get(f"/stocktake/sessions/{sid}/export.xlsx").content)).worksheets[0]
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith(("=", "+", "@"))]
    assert cells, "危険な値のセルが見つからない"
    assert all(c.data_type == "s" for c in cells), [(c.coordinate, c.data_type) for c in cells]


def test_G01_order_excel_writes_text_not_formula(admin_client):
    payload = {"vendor": "=evil", "ordered_at": "2026-10-23", "requester": {"department": "+x", "name": "@y"},
               "items": [{"item_name": DANGEROUS, "item_code": "-1", "quantity": 1, "unit_price": 1, "subtotal": 1, "notes": "=2"}],
               "total": 1, "has_unpriced_items": False}
    ws = load_workbook(io.BytesIO(admin_client.post("/orders/excel", json={"payloads": [payload]}).content)).worksheets[0]
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value[:1] in "=+-@"]
    assert len(cells) >= 4
    assert all(c.data_type == "s" for c in cells)


# ---------- G-02 アカウント単位のロック ----------
def _login(client, username, password):
    return client.post("/auth/login", json={"username": username, "password": password})


def test_G02_account_locks_after_10_failures(client, db):
    for _ in range(10):
        assert _login(client, "taro", "wrong").status_code == 401
    r = _login(client, "taro", "user1234")  # 正しくてもロック中は入れない
    assert r.status_code == 429
    assert "ロック" in r.json()["detail"]
    # 15分たてば入れる
    u = db.query(User).filter(User.username == "taro").one()
    db.refresh(u)
    u.locked_until = jst_now() - timedelta(seconds=1)
    db.commit()
    assert _login(client, "taro", "user1234").status_code == 200
    db.refresh(u)
    assert u.failed_login_count == 0


def test_G02_success_resets_failure_count(client, db):
    for _ in range(9):
        _login(client, "taro", "wrong")
    assert _login(client, "taro", "user1234").status_code == 200
    for _ in range(9):
        _login(client, "taro", "wrong")
    assert _login(client, "taro", "user1234").status_code == 200


def test_G02_admin_can_unlock(client, admin_client, db):
    for _ in range(10):
        _login(client, "taro", "wrong")
    taro = db.query(User).filter(User.username == "taro").one()
    users = {u["username"]: u for u in admin_client.get("/users").json()}
    assert users["taro"]["locked"] is True
    assert admin_client.patch(f"/users/{taro.id}", json={"unlock": True}).status_code == 200
    assert _login(client, "taro", "user1234").status_code == 200


# ---------- G-03 パスワード変更で他の端末のログインを無効に ----------
def test_G03_admin_reset_logs_out_existing_sessions(make_client, db):
    taro = make_client("taro")
    admin = make_client("admin")
    uid = db.query(User).filter(User.username == "taro").one().id
    assert admin.patch(f"/users/{uid}", json={"password": "resetpass1"}).status_code == 200
    assert taro.get("/auth/me").status_code == 401
    assert taro.post("/auth/refresh").status_code == 401  # リフレッシュでも戻れない


def test_G03_self_change_keeps_current_session_only(make_client):
    pc = make_client("taro")
    phone = make_client("taro")
    r = pc.post("/auth/password", json={"current_password": "user1234", "new_password": "brandnew99"})
    assert r.status_code == 200
    assert pc.get("/auth/me").status_code == 200  # 変えた本人の画面は続けられる
    assert phone.get("/auth/me").status_code == 401
    assert phone.post("/auth/refresh").status_code == 401


def test_G03_deactivate_revokes_refresh(make_client, db):
    taro = make_client("taro")
    admin = make_client("admin")
    uid = db.query(User).filter(User.username == "taro").one().id
    admin.delete(f"/users/{uid}")
    admin.patch(f"/users/{uid}", json={"is_active": 1})
    assert taro.post("/auth/refresh").status_code == 401


# ---------- G-04 ユーザーの有無で応答の重さを変えない ----------
def test_G04_unknown_user_still_runs_password_check(client, monkeypatch):
    import app.routers.auth as auth_router

    calls = []
    real = auth_router.verify_password
    monkeypatch.setattr(auth_router, "verify_password", lambda p, h: calls.append(h) or real(p, h))
    assert _login(client, "nobody-here", "whatever1").status_code == 401
    assert len(calls) == 1


# ---------- G-05 bcrypt の 72 バイト上限 ----------
def test_G05_password_over_72_bytes_rejected(admin_client):
    r = admin_client.post("/users", json={"username": "longpw", "password": "あ" * 25, "display_name": "x", "user_role": "USER"})
    assert r.status_code == 422
    assert "72" in r.json()["detail"]
    ok = admin_client.post("/users", json={"username": "longpw", "password": "a" * 72, "display_name": "x", "user_role": "USER"})
    assert ok.status_code == 201


# ---------- G-07 他サイトからの書き込みを拒否 ----------
def test_G07_cross_site_write_rejected(user_client, item_id):
    r = user_client.post("/outbound", json={"item_id": item_id, "quantity": 1}, headers={"Sec-Fetch-Site": "cross-site"})
    assert r.status_code == 403
    ok = user_client.post("/outbound", json={"item_id": item_id, "quantity": 1}, headers={"Sec-Fetch-Site": "same-origin"})
    assert ok.status_code == 201
    # 読み取りは止めない
    assert user_client.get("/inventory", headers={"Sec-Fetch-Site": "cross-site"}).status_code == 200


# ---------- G-09 本番設定の確認 ----------
@pytest.mark.parametrize("secret", ["change-me", "short-secret"])
def test_G09_production_rejects_weak_secret(secret):
    from app.config import Settings, validate_settings

    with pytest.raises(RuntimeError):
        validate_settings(Settings(app_env="production", jwt_secret=secret))


def test_G09_production_warns_without_secure_cookie(caplog):
    from app.config import Settings, validate_settings

    with caplog.at_level(logging.WARNING):
        validate_settings(Settings(app_env="production", jwt_secret="x" * 40, cookie_secure=False))
    assert "COOKIE_SECURE" in caplog.text


# ---------- G-11 セキュリティの記録 ----------
def test_G11_security_events_logged(client, caplog):
    with caplog.at_level(logging.INFO, logger="security"):
        for _ in range(10):
            _login(client, "taro", "wrong")
    text = caplog.text
    assert "login_failed" in text and "account_locked" in text and "taro" in text
    assert "wrong" not in text  # パスワードは残さない
